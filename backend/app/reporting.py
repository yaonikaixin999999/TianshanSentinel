from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .settings import settings


BLUE = "1769E8"
NAVY = "15304F"
MUTED = "6B7D93"
LIGHT_BLUE = "EAF3FF"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
CONTENT_WIDTH_DXA = 9360


def _runtime_path(url: str) -> Path:
    relative = url.removeprefix("/runtime/")
    path = (settings.runtime_dir / relative).resolve()
    runtime_root = settings.runtime_dir.resolve()
    if runtime_root not in path.parents or not path.is_file():
        raise FileNotFoundError(url)
    return path


def _set_run_font(run, size: float, color: str = NAVY, bold: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indentation = properties.find(qn("w:tblInd"))
    if indentation is None:
        indentation = OxmlElement("w:tblInd")
        properties.append(indentation)
    indentation.set(qn("w:w"), str(indent))
    indentation.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_width.set(qn("w:w"), str(widths[index]))
            tc_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _mark_header_row(table) -> None:
    properties = table.rows[0]._tr.get_or_add_trPr()
    marker = properties.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        properties.append(marker)
    marker.set(qn("w:val"), "true")


def _set_picture_alt(inline_shape, description: str) -> None:
    inline_shape._inline.docPr.set("descr", description)
    inline_shape._inline.docPr.set("title", description)


def _style_cell_text(cell, size: float = 9.5, color: str = NAVY, bold: bool = False) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        for run in paragraph.runs:
            _set_run_font(run, size, color, bold)


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)


def _event_summary(regions: list[dict[str, Any]]) -> str:
    labels: dict[str, int] = {}
    for region in regions:
        label = region.get("event_label", "待判定")
        labels[label] = labels.get(label, 0) + 1
    return "、".join(f"{label} {count} 处" for label, count in labels.items()) or "无有效变化斑块"


def build_geojson(analysis: dict[str, Any], project: dict[str, Any]) -> bytes:
    features = []
    for region in analysis["regions"]:
        x, y = region["x"], region["y"]
        x2, y2 = x + region["width"], y + region["height"]
        features.append(
            {
                "type": "Feature",
                "id": region["region_id"],
                "geometry": {
                    "type": "Polygon",
                    "coordinates": [[[x, y], [x2, y], [x2, y2], [x, y2], [x, y]]],
                },
                "properties": {
                    "analysis_id": analysis["id"],
                    "project_id": project["id"],
                    "project_name": project["name"],
                    "coordinate_system": "image_pixel",
                    "area_pixels": region["area_pixels"],
                    "area_ratio": region["area_ratio"],
                    "confidence": region["mean_confidence"],
                    "risk_level": region["risk_level"],
                    "event_type": region.get("event_type", "unclassified"),
                    "event_label": region.get("event_label", "待判定"),
                    "review_status": region.get("review_status", "pending"),
                    "reviewer_note": region.get("reviewer_note", ""),
                },
            }
        )
    collection = {
        "type": "FeatureCollection",
        "name": f"tianshan-sentinel-{analysis['id'][:8]}",
        "coordinateReferenceSystem": "image_pixel",
        "notice": "源影像未提供地理参考时，坐标为影像像素坐标。",
        "features": features,
    }
    return json.dumps(collection, ensure_ascii=False, indent=2).encode("utf-8")


def build_project_csv(timeline: dict[str, Any]) -> bytes:
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(["任务编号", "前时相", "后时相", "分析时间", "变化率", "置信度", "斑块数", "风险等级", "复核状态"])
    for item in timeline["items"]:
        writer.writerow(
            [
                item["id"], item["before_label"], item["after_label"], item["created_at"],
                item["change_ratio"], item["mean_confidence"], item["region_count"],
                item["overall_risk"], item["review_status"],
            ]
        )
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def build_docx_report(analysis: dict[str, Any], project: dict[str, Any]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    heading_tokens = {
        "Heading 1": (16, 16, 8, BLUE),
        "Heading 2": (13, 12, 6, BLUE),
        "Heading 3": (12, 8, 4, NAVY),
    }
    for style_name, (size, before, after, color) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_font(header.add_run("天山哨兵 | 遥感变化审计报告"), 9, MUTED, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(footer.add_run(f"任务 {analysis['id'][:8]} | 自动生成"), 8.5, MUTED)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    _set_run_font(kicker.add_run("CHANGE DETECTION AUDIT"), 9, BLUE, True)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    _set_run_font(title.add_run("遥感变化检测与研判报告"), 24, NAVY, True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    _set_run_font(subtitle.add_run(project["name"]), 13, MUTED)

    metadata = document.add_table(rows=5, cols=2)
    metadata.style = "Table Grid"
    metadata_rows = [
        ("任务编号", analysis["id"]),
        ("监测区域", project["area_name"]),
        ("时相范围", f"{analysis['before_label']}  至  {analysis['after_label']}"),
        ("分析时间", analysis["created_at"]),
        ("推理引擎", analysis["engine"]),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_rows):
        row.cells[0].text = label
        row.cells[1].text = str(value)
        _shade_cell(row.cells[0], LIGHT_BLUE)
        _style_cell_text(row.cells[0], 9.5, BLUE, True)
        _style_cell_text(row.cells[1], 9.5, NAVY)
    _set_table_geometry(metadata, [1800, 7560])
    _mark_header_row(metadata)

    _add_heading(document, "1. 核心结论")
    conclusion = document.add_paragraph(analysis["narrative"])
    conclusion.paragraph_format.line_spacing = 1.25
    conclusion.add_run(f" 事件类型建议：{_event_summary(analysis['regions'])}。")

    metrics = document.add_table(rows=2, cols=4)
    metrics.style = "Table Grid"
    labels = ["变化占比", "平均置信度", "有效斑块", "综合风险"]
    values = [
        f"{analysis['change_ratio']:.2%}",
        f"{analysis['mean_confidence']:.1%}",
        str(analysis["region_count"]),
        {"high": "高风险", "medium": "中风险", "low": "低风险"}.get(analysis["overall_risk"], analysis["overall_risk"]),
    ]
    for index, label in enumerate(labels):
        metrics.cell(0, index).text = label
        metrics.cell(1, index).text = values[index]
        _shade_cell(metrics.cell(0, index), LIGHT_BLUE)
        _style_cell_text(metrics.cell(0, index), 9, BLUE, True)
        _style_cell_text(metrics.cell(1, index), 13, NAVY, True)
        metrics.cell(0, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        metrics.cell(1, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_table_geometry(metrics, [2340, 2340, 2340, 2340])
    _mark_header_row(metrics)

    _add_heading(document, "2. 变化证据")
    overlay_path = _runtime_path(analysis["overlay_url"])
    overlay_picture = document.add_picture(str(overlay_path), width=Inches(6.35))
    _set_picture_alt(overlay_picture, "变化区域叠加结果，红色区域表示模型检测到的变化")
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph("图 1  变化区域叠加结果（红色区域为模型检测结果）")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    _set_run_font(caption.runs[0], 9, MUTED)

    comparison = document.add_table(rows=1, cols=2)
    comparison.autofit = False
    for cell, url, label in zip(
        comparison.rows[0].cells,
        (analysis["before_url"], analysis["after_url"]),
        (analysis["before_label"], analysis["after_label"]),
    ):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        phase_picture = paragraph.add_run().add_picture(str(_runtime_path(url)), width=Inches(3.02))
        _set_picture_alt(phase_picture, f"{label}遥感影像")
        label_paragraph = cell.add_paragraph(label)
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(label_paragraph.runs[0], 9, MUTED, True)
    _set_table_geometry(comparison, [4680, 4680])
    _mark_header_row(comparison)

    _add_heading(document, "3. 变化斑块清单")
    if analysis["regions"]:
        region_table = document.add_table(rows=1, cols=6)
        region_table.style = "Table Grid"
        headers = ["编号", "事件建议", "像素面积", "占比", "置信度", "复核状态"]
        for index, label in enumerate(headers):
            region_table.cell(0, index).text = label
            _shade_cell(region_table.cell(0, index), LIGHT_BLUE)
            _style_cell_text(region_table.cell(0, index), 8.5, BLUE, True)
        for region in analysis["regions"][:12]:
            row = region_table.add_row().cells
            values = [
                region["region_id"], region.get("event_label", "待判定"), region["area_pixels"],
                f"{region['area_ratio']:.2%}", f"{region['mean_confidence']:.1%}",
                {"pending": "待复核", "approved": "已确认", "rejected": "已排除"}.get(region.get("review_status"), "待复核"),
            ]
            for index, value in enumerate(values):
                row[index].text = str(value)
                _style_cell_text(row[index], 8.5, NAVY)
        _set_table_geometry(region_table, [650, 2500, 1400, 1300, 1450, 2060])
        _mark_header_row(region_table)
    else:
        document.add_paragraph("本次分析未检测到达到面积阈值的有效变化斑块。")

    _add_heading(document, "4. 方法与使用限制")
    notes = [
        "变化掩膜由当前部署的双时相变化检测模型生成，风险等级依据面积与置信度综合计算。",
        "新增、拆除和地表改造属于规则型辅助建议，必须经过人工复核后才能作为业务结论。",
        "当前 GeoJSON 默认使用影像像素坐标；只有带完整地理参考信息的源影像才能换算真实经纬度和面积。",
    ]
    for note in notes:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(note)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.167

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_pdf_report(analysis: dict[str, Any], project: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Image as PdfImage
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except KeyError:
        pass
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        leftMargin=0.72 * inch,
        rightMargin=0.72 * inch,
        topMargin=0.72 * inch,
        bottomMargin=0.68 * inch,
        title=f"天山哨兵变化检测报告-{analysis['id'][:8]}",
        author="天山哨兵",
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle("CNBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=9.5, leading=15, textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=7)
    h1 = ParagraphStyle("CNH1", parent=body, fontSize=14, leading=19, textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=14, spaceAfter=8)
    title_style = ParagraphStyle("CNTitle", parent=body, fontSize=23, leading=29, textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=6)
    kicker = ParagraphStyle("CNKicker", parent=body, fontSize=8.5, textColor=colors.HexColor(f"#{BLUE}"), spaceAfter=5)
    caption = ParagraphStyle("CNCaption", parent=body, fontSize=8, textColor=colors.HexColor(f"#{MUTED}"), alignment=TA_CENTER, spaceAfter=8)

    story = [
        Paragraph("CHANGE DETECTION AUDIT", kicker),
        Paragraph("遥感变化检测与研判报告", title_style),
        Paragraph(project["name"], ParagraphStyle("subtitle", parent=body, fontSize=12, textColor=colors.HexColor(f"#{MUTED}"), spaceAfter=16)),
    ]
    metadata = [
        ["任务编号", analysis["id"]],
        ["监测区域", project["area_name"]],
        ["时相范围", f"{analysis['before_label']} 至 {analysis['after_label']}"],
        ["分析时间", analysis["created_at"]],
        ["推理引擎", analysis["engine"]],
    ]
    meta_table = Table(metadata, colWidths=[1.25 * inch, 5.75 * inch], hAlign="LEFT")
    meta_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(f"#{NAVY}")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{LIGHT_BLUE}")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor(f"#{BLUE}")),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#D9E4F0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([meta_table, Paragraph("1. 核心结论", h1), Paragraph(analysis["narrative"] + f" 事件类型建议：{_event_summary(analysis['regions'])}。", body)])

    risk_label = {"high": "高风险", "medium": "中风险", "low": "低风险"}.get(analysis["overall_risk"], analysis["overall_risk"])
    metric_data = [
        ["变化占比", "平均置信度", "有效斑块", "综合风险"],
        [f"{analysis['change_ratio']:.2%}", f"{analysis['mean_confidence']:.1%}", str(analysis["region_count"]), risk_label],
    ]
    metric_table = Table(metric_data, colWidths=[1.75 * inch] * 4)
    metric_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{LIGHT_BLUE}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
        ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor(f"#{NAVY}")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#D9E4F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([metric_table, Paragraph("2. 变化证据", h1)])

    overlay_image = PdfImage(str(_runtime_path(analysis["overlay_url"])))
    overlay_image._restrictSize(7 * inch, 4.2 * inch)
    story.extend([overlay_image, Paragraph("图 1  变化区域叠加结果（红色区域为模型检测结果）", caption), Paragraph("3. 变化斑块清单", h1)])

    if analysis["regions"]:
        rows = [["编号", "事件建议", "像素面积", "占比", "置信度", "复核状态"]]
        for region in analysis["regions"][:15]:
            rows.append([
                str(region["region_id"]), region.get("event_label", "待判定"), str(region["area_pixels"]),
                f"{region['area_ratio']:.2%}", f"{region['mean_confidence']:.1%}",
                {"pending": "待复核", "approved": "已确认", "rejected": "已排除"}.get(region.get("review_status"), "待复核"),
            ])
        region_table = Table(rows, colWidths=[0.55 * inch, 1.8 * inch, 1.05 * inch, 0.9 * inch, 1.05 * inch, 1.4 * inch], repeatRows=1)
        region_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{LIGHT_BLUE}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
            ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor(f"#{NAVY}")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9E4F0")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(region_table)
    else:
        story.append(Paragraph("本次分析未检测到达到面积阈值的有效变化斑块。", body))

    story.extend([
        Paragraph("4. 方法与使用限制", h1),
        Paragraph("1. 变化掩膜由当前部署的双时相变化检测模型生成，风险等级依据面积与置信度综合计算。", body),
        Paragraph("2. 新增、拆除和地表改造属于规则型辅助建议，必须经过人工复核后才能作为业务结论。", body),
        Paragraph("3. 当前 GeoJSON 默认使用影像像素坐标；只有带完整地理参考信息的源影像才能换算真实经纬度和面积。", body),
    ])

    def add_page_furniture(canvas, doc):
        canvas.saveState()
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawString(0.72 * inch, 10.55 * inch, "天山哨兵 | 遥感变化审计报告")
        canvas.drawRightString(7.78 * inch, 0.38 * inch, f"第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=add_page_furniture, onLaterPages=add_page_furniture)
    return output.getvalue()
